#!/usr/bin/env python
"""Seed loader for the NOVA 2.0 RCA demo.

Reads the sample workbook (``data/Sample_Log_Data.xlsx``) and loads its three
sheets into the ``audit_transaction`` / ``audit_error`` / ``audit_service``
collections, applying ONLY the repairs required to undo *Excel-inflicted*
damage (see brief §2.5). The authentic messiness of the source system is
preserved verbatim so the agent has to cope with it:

  * ``usr_id`` is float-mangled by Excel (``1.73070901334081e+18``). We recover
    the true 19-digit value from a payload where present, else synthesize a
    deterministic 19-digit id. Stored as a STRING.
  * literal ``'NULL'`` strings are KEPT in ``biz_err_cd_*`` / ``biz_err_cd_*_msg``
    / ``trace_id`` / ``nfrchk`` (authentic system behaviour); they are only
    treated as "absent" for the mangled-float field (``usr_id``).
  * mojibake ``â€¢`` is repaired to ``•``.
  * timestamps (``2026-07-07 16:09:38.022+05:30``, tolerating truncated ms like
    ``16:15:20.59``) are parsed to UTC BSON ``Date``.
  * a monotonic ``seq`` (int) is added per ``trns_id``.
  * ``channel`` (mobile/web) is derived from ``chnl_id`` on transactions.

Spelling drift (``trns_stts`` vs ``trns_sts``), the ``' UNIPAY'`` leading space
and double-encoded JSON are all preserved.

Idempotent: the three collections are cleared (``delete_many({})``) before
insert. With ``--kb`` the curated ``kb_docs`` knowledge base is also seeded.

Search / vector indexes are NEVER touched here.

Usage::

    . .venv/bin/activate
    python data/seed.py            # load the 3 audit collections
    python data/seed.py --kb       # also seed kb_docs
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import re
import sys
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = REPO_ROOT / "data" / "Sample_Log_Data.xlsx"

# Reuse the project's single source of truth for connection + collection names.
sys.path.insert(0, str(REPO_ROOT / "agent"))
from nova_agent.config import settings  # noqa: E402
from nova_agent.db import get_db  # noqa: E402

# Sheet -> collection.  Note the trailing space in "audit error " is authentic.
SHEETS = {
    "audit transaction": settings.C_TXN,
    "audit error ": settings.C_ERR,
    "audit service": settings.C_SVC,
}

BANNER_ROW = 1
HEADER_ROW = 3
DATA_START_ROW = 4

# ---------------------------------------------------------------------------
# Repair helpers
# ---------------------------------------------------------------------------

_MOJIBAKE = {"â€¢": "•"}


def repair_mojibake(value):
    if isinstance(value, str):
        for bad, good in _MOJIBAKE.items():
            if bad in value:
                value = value.replace(bad, good)
    return value


def parse_ts(value):
    """Parse a workbook timestamp to a tz-aware UTC datetime (BSON Date)."""
    if value is None or value == "" or value == "NULL":
        return None
    if isinstance(value, dt.datetime):
        d = value
        if d.tzinfo is None:
            d = d.replace(tzinfo=dt.timezone.utc)
        return d.astimezone(dt.timezone.utc)
    s = str(value).strip()
    # datetime.fromisoformat (3.11+) tolerates the space separator, the
    # +05:30 offset and fractional seconds of any 1-6 digit length.
    d = dt.datetime.fromisoformat(s)
    if d.tzinfo is None:
        d = d.replace(tzinfo=dt.timezone.utc)
    return d.astimezone(dt.timezone.utc)


def _sig_digits(d: float) -> int:
    s = "%.15g" % float(d)
    mantissa = s.split("e")[0].replace("-", "").replace(".", "").lstrip("0")
    return max(1, len(mantissa))


def _round_sig(x, k: int) -> str:
    return "%.{}g".format(k) % float(x)


def build_usr_id_recovery(rows_by_sheet: dict[str, list[dict]]):
    """Learn true 19-digit usr_ids from payloads, keyed by the mangled double.

    Excel truncates ``usr_id`` to ~15 significant figures on the transaction /
    error sheets and to ~6 on the service sheet. A single true value therefore
    shows up as *different* mangled doubles. We recover it once from the
    payloads (where a ``usrId`` / ``usr_id`` key still carries the full integer)
    and reuse it for every mangled double that is a coarser rounding of it — so
    the same customer stays the same id across all three collections.
    """
    # Candidate true ids: 19-digit ints written against a usr id key.
    pool: set[str] = set()
    key_pat = re.compile(r'"usr_?[iI]d"\s*:\s*"?(\d{19})"?')
    payload_fields = (
        "biz_req_input",
        "biz_resp_output",
        "err_dtls",
        "srvc_rqst",
        "srvc_rspn",
    )
    for rows in rows_by_sheet.values():
        for row in rows:
            for f in payload_fields:
                v = row.get(f)
                if isinstance(v, str):
                    pool.update(key_pat.findall(v))

    cache: dict[float, str] = {}

    def resolve(raw):
        if raw is None:
            return None
        if isinstance(raw, str):
            return None if raw == "NULL" else raw
        d = float(raw)
        if d in cache:
            return cache[d]
        k = _sig_digits(d)
        recovered = None
        for cand in pool:
            if _round_sig(cand, k) == _round_sig(d, k):
                recovered = cand
                break
        if recovered is None:
            # Deterministic 19-digit synthesis, stable per mangled double.
            h = hashlib.sha1(repr(d).encode()).hexdigest()
            recovered = str(10**18 + int(h, 16) % (10**18))
        cache[d] = recovered
        return recovered

    return resolve


def _as_str(value):
    """Coerce an Excel-numeric-coerced id back to a string (contract wants str)."""
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


CHANNEL = {2001: "mobile", 1001: "web"}


# ---------------------------------------------------------------------------
# Row -> document builders
# ---------------------------------------------------------------------------


def build_txn(row: dict, usr_id: str) -> dict:
    return {
        "msg_id": int(row["msg_id"]),
        "trns_id": str(row["trns_id"]),
        "sessn_tkn_id": str(row["sessn_tkn_id"]),
        "clnt_ip_addr": str(row["clnt_ip_addr"]),
        "req_no": _as_str(row["req_no"]),
        "usr_rltnshp_no": _as_str(row["usr_rltnshp_no"]),
        "trns_stts": str(row["trns_stts"]),
        "strt_time": parse_ts(row["strt_time"]),
        "end_time": parse_ts(row["end_time"]),
        "usr_id": usr_id,
        "usr_typ": str(row["usr_typ"]),
        "cmnd_id": str(row["cmnd_id"]),
        "chnl_id": int(row["chnl_id"]),
        "biz_req_input": repair_mojibake(row["biz_req_input"]),
        "biz_resp_output": repair_mojibake(row["biz_resp_output"]),
        "crtd_by": str(row["crtd_by"]),
        "crtd_on": parse_ts(row["crtd_on"]),
        "trace_id": row["trace_id"],  # keep 'NULL' or hex verbatim
        "nfrchk": row["nfrchk"],  # keep 'NULL'
        "channel": CHANNEL.get(int(row["chnl_id"]), "mobile"),
    }


def build_err(row: dict, usr_id: str) -> dict:
    doc = {
        "msg_id": int(row["msg_id"]),
        "trns_id": str(row["trns_id"]),
        "sessn_tkn_id": str(row["sessn_tkn_id"]),
        "err_cd": _as_str(row["err_cd"]),  # 5151 -> "5151"; BS* kept
        "err_typ": str(row["err_typ"]),
        "err_dscrptn": repair_mojibake(row["err_dscrptn"]),
        "err_dtls": repair_mojibake(row["err_dtls"]),  # JSON string kept verbatim
        "crtd_by": str(row["crtd_by"]),
        "crtd_on": parse_ts(row["crtd_on"]),
        "err_time": parse_ts(row["err_time"]),
        "usr_id": usr_id,
    }
    for i in range(1, 6):
        doc[f"biz_err_cd_{i}"] = row.get(f"biz_err_cd_{i}")  # keep 'NULL'
        doc[f"biz_err_cd_{i}_msg"] = repair_mojibake(row.get(f"biz_err_cd_{i}_msg"))
    return doc


def build_svc(row: dict, usr_id: str) -> dict:
    err_cd = row["err_cd"]
    if err_cd is None:
        err_cd = ""  # contract: str | int | ""
    elif isinstance(err_cd, float) and err_cd.is_integer():
        err_cd = int(err_cd)
    err_msg = row["err_msg"]
    if err_msg is None:
        err_msg = ""
    trns_sts = row["trns_sts"]
    if isinstance(trns_sts, float) and trns_sts.is_integer():
        trns_sts = int(trns_sts)
    return {
        "msg_id": int(row["msg_id"]),
        "trns_id": str(row["trns_id"]),
        "sessn_tkn_id": str(row["sessn_tkn_id"]),
        "trns_sts": str(trns_sts),  # 'success' | '0' | '-2'
        "err_cd": err_cd,
        "err_msg": repair_mojibake(err_msg),
        "usr_id": usr_id,
        "srvc_rqst": repair_mojibake(row["srvc_rqst"]),
        "srvc_rspn": repair_mojibake(row["srvc_rspn"]),
        "srvc_id": str(row["srvc_id"]),
        "systm_nm": row["systm_nm"],  # preserve ' UNIPAY' leading space
        "api_url": str(row["api_url"]),
        "rqst_id_sent": _as_str(row["rqst_id_sent"]),
        "rqst_id_rcvd": _as_str(row["rqst_id_rcvd"]),
        "crtd_by": str(row["crtd_by"]),
        "crtd_on": parse_ts(row["crtd_on"]),
        "strt_time": parse_ts(row["strt_time"]),
        "end_time": parse_ts(row["end_time"]),
    }


def assign_seq(docs: list[dict], time_field: str) -> None:
    """Add a monotonic per-trns_id seq (1-based), ordered by time then msg_id."""
    groups: dict[str, list[dict]] = {}
    for d in docs:
        groups.setdefault(d["trns_id"], []).append(d)
    epoch = dt.datetime(1970, 1, 1, tzinfo=dt.timezone.utc)
    for group in groups.values():
        group.sort(key=lambda d: (d.get(time_field) or epoch, d["msg_id"]))
        for i, d in enumerate(group, start=1):
            d["seq"] = i


# ---------------------------------------------------------------------------
# Workbook reading
# ---------------------------------------------------------------------------


def read_sheet(wb, sheet_name: str) -> list[dict]:
    ws = wb[sheet_name]
    headers = [ws.cell(row=HEADER_ROW, column=c).value for c in range(1, ws.max_column + 1)]
    rows: list[dict] = []
    for r in range(DATA_START_ROW, ws.max_row + 1):
        vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        if all(v is None for v in vals):
            continue
        row = {headers[i]: vals[i] for i in range(len(headers)) if headers[i] is not None}
        rows.append(row)
    return rows


def load_audit(db) -> dict[str, int]:
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True, read_only=True)
    raw = {sheet: read_sheet(wb, sheet) for sheet in SHEETS}
    resolve_usr_id = build_usr_id_recovery(raw)

    txn_docs = [build_txn(r, resolve_usr_id(r.get("usr_id"))) for r in raw["audit transaction"]]
    err_docs = [build_err(r, resolve_usr_id(r.get("usr_id"))) for r in raw["audit error "]]
    svc_docs = [build_svc(r, resolve_usr_id(r.get("usr_id"))) for r in raw["audit service"]]

    assign_seq(txn_docs, "strt_time")
    assign_seq(err_docs, "err_time")
    assign_seq(svc_docs, "strt_time")

    counts = {}
    for coll_name, docs in (
        (settings.C_TXN, txn_docs),
        (settings.C_ERR, err_docs),
        (settings.C_SVC, svc_docs),
    ):
        coll = db[coll_name]
        coll.delete_many({})
        if docs:
            coll.insert_many(docs)
        counts[coll_name] = len(docs)
    return counts


# ---------------------------------------------------------------------------
# Knowledge base (kb_docs) — ~15 curated runbook / ownership entries
# ---------------------------------------------------------------------------

KB_DOCS = [
    {
        "title": "EIS gateway route ownership map",
        "body": (
            "The Enterprise Integration Services (EIS) gateway fronts every "
            "surrounding-system call from NOVA 2.0. Routes are namespaced by "
            "destination under the gateway path. The thirdParty/wrapper family "
            "carries external-vendor traffic: thirdParty/wrapper/Martech routes "
            "to the Martech (MARTECH_DWH) offer engine that powers "
            "eis_marTech_getInteractOffers; thirdParty/wrapper/services carries "
            "THINK360 (IT_JVP) financial-fitness token and summary calls. The "
            "misc/microservice routes (e.g. CustEnquiry) carry EIS-native "
            "customer-profile lookups. Ownership: the EIS platform team owns the "
            "gateway, connection pools and TLS; the destination vendor owns the "
            "business logic behind each wrapper route. When many distinct srvc_id "
            "values fail at once, suspect the shared gateway before any single "
            "vendor."
        ),
        "codes": ["TEEIS009", "5151"],
        "tags": ["eis", "gateway", "routing", "ownership", "thirdParty", "wrapper", "martech"],
    },
    {
        "title": "Runbook: TEEIS009 Response Transformation failed",
        "body": (
            "TEEIS009 'Response Transformation failed.' is emitted by the EIS "
            "gateway when a downstream vendor returns a payload the EIS response "
            "transformer cannot map to the expected schema — typically because "
            "the vendor is slow, returning truncated/garbled bodies, or returning "
            "an unexpected error envelope. On audit_service it appears with "
            "trns_sts -2 and a long duration (end_time - strt_time in the "
            "multi-second range, e.g. ~10s). The matching journey-side failure is "
            "usually a generic BS000001 business exception (for offers, "
            "BSYHP003). Triage: (1) confirm incidence — count distinct usr_id / "
            "sessn_tkn_id with TEEIS009 on the same srvc_id in the window; many "
            "users on one route means a downstream brownout, one user means a "
            "client problem. (2) Compare p95 duration for the route against the "
            "30-minute baseline in service_metrics. (3) Name the shared EIS route "
            "(e.g. thirdParty/wrapper/Martech) as origin, not the loud journey "
            "error. Immediate mitigation: circuit-break the affected widget; fix: "
            "engage the vendor / raise EIS timeout; prevention: alert on route "
            "p95 and TEEIS009 rate."
        ),
        "codes": ["TEEIS009", "BS000001", "BSYHP003"],
        "tags": ["runbook", "eis", "transformation", "brownout", "martech"],
    },
    {
        "title": "UNIPAY BBPS compliance codes and layered status",
        "body": (
            "UNIPAY (systm_nm ' UNIPAY', note the authentic leading space) is the "
            "BBPS autopay/billpay provider. Its responses carry a layered status: "
            "the transport status audit_service.trns_sts can read 'success' while "
            "the JSON srvc_rspn reason block reports a business failure. Decode "
            "reason.responseCode: '200' with responseReason 'No data found' is a "
            "benign empty result; '001' with responseReason 'Failure' is a real "
            "business failure and carries a complianceReason such as BFR003 (no "
            "bill data / biller failure) or complianceRespCd SIP019 (standing "
            "instruction not available). A rising ratio of responseCode 001 / "
            "BFR003 over minutes while trns_sts stays 'success' is a UNIPAY BBPS "
            "brownout that never turns a status field red — it is only visible by "
            "decoding the payload and trending business-failure rate in "
            "service_metrics."
        ),
        "codes": ["BFR003", "SIP019", "001", "200"],
        "tags": ["unipay", "bbps", "compliance", "layered-status", "billpay", "autopay"],
    },
    {
        "title": "THINK360 (IT_JVP) token lifecycle and 5151",
        "body": (
            "PFM financial-fitness journeys (pfm_financial-fitness_customer_kyc_"
            "fetch, _summary_fetch) are backed by THINK360, reached through EIS "
            "destination IT_JVP. The flow first calls "
            "eis_financialfitness_accessToken (TXN_TYPE THINK360, TXN_SUB_TYPE "
            "LOGIN) which returns an EIS_RESPONSE with a JWT access_token and an "
            "'expiry' epoch field; subsequent summary/kyc calls send that token as "
            "a Bearer header. When the token renewal fails or the JWT 'expiry' has "
            "passed, the dependent PFM calls fail almost instantly (single-digit "
            "ms) with legacy technical err_cd 5151 ('Exception occurred while "
            "executing the service'), or with a 404 'Data Not Found' buried inside "
            "an EIS_RESPONSE while trns_sts reads 0. The credential call itself "
            "barely registers, so the origin is the shared token dependency — a "
            "fan-out from one credential, not many independent failures."
        ),
        "codes": ["5151", "TEEIS009"],
        "tags": ["think360", "it_jvp", "token", "jwt", "pfm", "financial-fitness", "credential"],
    },
    {
        "title": "Runbook: unwrapping nested business error chains (BSDML050 -> BSDML024)",
        "body": (
            "audit_error.err_dtls is a JSON string keyed by business error code "
            "and can nest. The public/outer code is a customer-facing wrapper; the "
            "innermost code names the real dependency failure. For demand-loan "
            "closure (lending_demand_demand-loan-closure_manage) the outer "
            "BSDML050 ('Your Demand Loan closure request has failed. Kindly visit "
            "branch...') wraps an inner BSDML024 whose data carries "
            "loan_discharge: 'received improper response'. That inner text means "
            "the CBS loan-discharge dependency returned a malformed / improper "
            "body — a contract violation, not a NOVA bug. Always unwrap err_dtls to "
            "the innermost code before naming a root cause; report both the "
            "public code and the innermost cause."
        ),
        "codes": ["BSDML050", "BSDML024"],
        "tags": ["runbook", "nested-error", "unwrap", "lending", "demand-loan", "cbs"],
    },
    {
        "title": "Martech CIF desync — BSYHP003 CIF Is Not Recognized",
        "body": (
            "BSYHP003 'CIF Is Not Recognized By Martech.' surfaces on "
            "martech_hyper-personalization_offers_fetch as a BUSINESS_EXCEPTION. "
            "It means the customer's CIF is not present in the Martech audience "
            "store. A single BSYHP003 is normal (new/edge customer); a spike for a "
            "cohort usually follows a Martech audience refresh — the ops "
            "transaction ops_martech_audience_refresh_execute, which is logged as "
            "a SUCCESS document, not an error. When a refresh desyncs or drops "
            "CIFs, offers fetch starts returning BSYHP003 for the affected cohort "
            "minutes later. Triage: look back ~10 minutes for a Martech audience "
            "refresh ops transaction preceding the spike; the non-error precursor "
            "is the origin. This is distinct from eis-martech-brownout, where "
            "BSYHP003/BS000001 ride on top of TEEIS009 gateway transformation "
            "failures."
        ),
        "codes": ["BSYHP003", "BS000001"],
        "tags": ["martech", "cif", "desync", "audience-refresh", "precursor", "ops"],
    },
    {
        "title": "GSA insurance validation timeout — TEY2GSA02",
        "body": (
            "gi_gsa_purchase_validate_checks (GI insurance purchase, web channel "
            "chnl_id 1001) validates a purchase against an external insurer "
            "dependency. When that dependency slows, the journey hangs (p95 climbs "
            "toward ~15s) and then throws TEY2GSA02, a TECHNICAL_EXCEPTION whose "
            "description mentions request validation / processing failing 'due to "
            "unexpected data state'. The err_cd is timeout-shaped: the technical "
            "exception is a symptom of the dependency's latency, not a defect in "
            "the validation code itself. Root cause is the slow insurer "
            "validation dependency; prove it with the duration (end_time - "
            "strt_time) and the p95 baseline comparison, not the error text."
        ),
        "codes": ["TEY2GSA02"],
        "tags": ["gsa", "insurance", "gi", "timeout", "validation", "web"],
    },
    {
        "title": "Runbook: EIS gateway connection-pool exhaustion",
        "body": (
            "The EIS gateway maintains a bounded connection pool per host. Under "
            "load or when a downstream stalls holding connections, the pool "
            "exhausts and NEW connections to the EIS gateway host fail across "
            "ALL eis_* routes simultaneously — Martech offers, financial-fitness "
            "token/summary, customer-profile CustEnquiry and CC account summary "
            "all degrade at once. This is the widest blast radius signature: the "
            "loudest errors are per-journey business exceptions (BS000001, "
            "BSYHP003) on many different journeys, but the single shared origin is "
            "the EIS gateway host/pool. Distinguish from a single-route brownout "
            "by counting how many distinct srvc_id / systm_nm EIS routes fail in "
            "the same window: many routes at once => gateway/pool, one route => "
            "that vendor. Immediate: shed load / raise pool size; prevention: "
            "alert on gateway connection-error rate and pool saturation."
        ),
        "codes": ["TEEIS009", "5151", "BS000001"],
        "tags": ["runbook", "eis", "gateway", "pool", "exhaustion", "blast-radius"],
    },
    {
        "title": "err_cd vocabulary glossary (BS* / TE* / numeric)",
        "body": (
            "NOVA error codes span three vocabularies. BS* are BUSINESS_"
            "EXCEPTION codes: BS000001 is the generic business exception wrapper; "
            "specific business codes appear inside err_dtls and biz_err_cd_* slots "
            "(BSYHP003 Martech CIF, BSPAPL010 PAPL eligibility, BSDML050/BSDML024 "
            "demand-loan). TE* are TECHNICAL_EXCEPTION codes, often journey- or "
            "system-specific: TEEIS009 EIS response transformation failed, "
            "TEY2GSA02 GSA validation/processing. Bare numerics are legacy "
            "technical codes: 5151 is a generic 'exception occurred while "
            "executing the service'. On audit_service err_cd also carries "
            "transport-ish numerics: 200 (HTTP-ish ok / no data), 1 (internal "
            "failure marker), 0 and -2 as status markers. Never treat a BS000001 "
            "as the root cause — it is a wrapper; unwrap to the specific code."
        ),
        "codes": ["BS000001", "BSYHP003", "BSPAPL010", "BSDML050", "BSDML024", "TEEIS009", "TEY2GSA02", "5151"],
        "tags": ["glossary", "err_cd", "business", "technical", "legacy", "vocabulary"],
    },
    {
        "title": "PAPL eligibility — BSPAPL010",
        "body": (
            "BSPAPL010 'Sorry! This offer is not available to you.' is a "
            "BUSINESS_EXCEPTION on loan_personal_papl_getEligibleLoanDetails "
            "(pre-approved personal loan eligibility). It is an expected, "
            "customer-specific outcome: the customer is simply not eligible for a "
            "PAPL offer at this time. A low background rate of BSPAPL010 is noise "
            "and must NOT be escalated as an incident. Only a sudden cohort-wide "
            "spike (which would indicate an eligibility-service or rules problem) "
            "is noteworthy; otherwise treat as benign business noise."
        ),
        "codes": ["BSPAPL010"],
        "tags": ["papl", "lending", "eligibility", "noise", "benign"],
    },
    {
        "title": "Layered status semantics — transport vs business",
        "body": (
            "Status fields in the audit collections lie by design and must be "
            "decoded on two axes. Transport status: audit_transaction.trns_stts "
            "('error'/'success', double-t) and audit_service.trns_sts "
            "('success'/'0'/'-2', single-t) describe whether the call completed at "
            "the wire level. Business status lives INSIDE the payload: UNIPAY "
            "srvc_rspn.reason.responseCode ('001'/'Failure' + complianceReason) "
            "and EIS EIS_RESPONSE.success/status_code (e.g. success:false, "
            "status_code 404). A call can be transport-success but business-fail "
            "(UNIPAY BFR003) or trns_sts 0 but carry a 404 (THINK360 summary). "
            "The correlate step must JSON-parse payloads up to two levels (they "
            "are often double-encoded) and classify each event on BOTH axes; "
            "trusting the top-level status yields the wrong RCA."
        ),
        "codes": ["BFR003", "TEEIS009"],
        "tags": ["layered-status", "transport", "business", "decode", "correlate"],
    },
    {
        "title": "Correlation keys and the widening ladder",
        "body": (
            "The correlation spine is sessn_tkn_id (a user session) -> trns_id "
            "(one journey action / API call) -> the three audit collections "
            "joined on trns_id. trace_id is unreliable (mostly the literal string "
            "'NULL'); trns_id is the join key. Investigation widens along a "
            "ladder: start at the failing trns_id, add its session context, then "
            "find the same cmnd_id (journey) or srvc_id (downstream route) across "
            "OTHER sessions in the +/-120s window, then compare against the "
            "service_metrics baseline (current window vs the preceding 30 "
            "minutes). Incidence across many sessions/users turns 'one user hit an "
            "error' into 'the route is down for everyone' — and separates a "
            "systemic outage from a single-client problem."
        ),
        "codes": [],
        "tags": ["correlation", "trns_id", "sessn_tkn_id", "incidence", "collect", "baseline"],
    },
    {
        "title": "Poison-retry loop vs systemic outage",
        "body": (
            "A single client app that replays a malformed biz_req_input produces "
            "identical technical exceptions (commonly legacy 5151) repeating for "
            "ONE usr_id / sessn_tkn_id. This looks alarming by raw error count but "
            "is NOT a systemic outage. The discriminator is incidence: count "
            "distinct usr_id and sessn_tkn_id carrying the signature. Many "
            "distinct users => systemic (route/gateway/dependency); one user or "
            "session repeating => a client-side poison-retry loop. Mitigation "
            "differs entirely: rate-limit / block the offending client versus "
            "engaging a downstream owner. Always run the incidence check before "
            "declaring an outage."
        ),
        "codes": ["5151"],
        "tags": ["poison-retry", "client", "incidence", "5151", "single-user"],
    },
    {
        "title": "service_metrics baseline and time-series trending",
        "body": (
            "service_metrics is a time-series collection (ts / meta) recording per "
            "downstream and per journey: durationMs, transportOk and businessOk, "
            "with meta.service (srvc_id or cmnd_id), meta.system, meta.kind "
            "(journey|downstream) and meta.route. It is the baseline oracle: "
            "compute p50/p95 duration and error/business-failure rate for the "
            "incident window and compare against the preceding 30-minute baseline "
            "for the same route. Latency-shaped incidents (EIS brownout, GSA hang, "
            "gateway pool exhaustion) show a p95 blowout; layered-business "
            "brownouts (UNIPAY BFR003) show a rising businessOk-false rate while "
            "transportOk stays true. Threshold alerts on these metrics are the "
            "usual 'prevention' recommendation."
        ),
        "codes": [],
        "tags": ["service_metrics", "time-series", "baseline", "p95", "trend", "prevention"],
    },
    {
        "title": "EIS payload envelopes and route destinations",
        "body": (
            "EIS calls (systm_nm 'EIS') use TXN_TYPE to select a destination. "
            "TXN_TYPE INTERACT / DESTINATION MARTECH_DWH is the Martech offer "
            "engine (eis_marTech_getInteractOffers) over the "
            "thirdParty/wrapper/Martech route; the request carries an EIS_PAYLOAD "
            "with getOffersForMultipleInteractionPoints. TXN_TYPE THINK360 / "
            "DESTINATION IT_JVP is financial-fitness: TXN_SUB_TYPE LOGIN returns "
            "the access token, FFS_FETCH_SUMMARY fetches the summary. EIS-native "
            "microservice calls (querycustprofile / CustEnquiry, "
            "cifAssociatedAccount) carry a plain EIS_PAYLOAD. Responses come as "
            "either {ERROR_CODE, EIS_RESPONSE:{status_code,success,message}, "
            "RESPONSE_STATUS} or {TXN_TYPE, EIS_PAYLOAD}. Reading TXN_TYPE / "
            "DESTINATION tells you which vendor sits behind a failing EIS route."
        ),
        "codes": ["TEEIS009"],
        "tags": ["eis", "payload", "envelope", "martech", "think360", "destination"],
    },
]


def load_kb(db) -> int:
    coll = db[settings.C_KB]
    coll.delete_many({})
    coll.insert_many([dict(d) for d in KB_DOCS])
    return len(KB_DOCS)


# ---------------------------------------------------------------------------
# main
# ---------------------------------------------------------------------------


def main() -> None:
    parser = argparse.ArgumentParser(description="Seed the NOVA RCA demo collections.")
    parser.add_argument("--kb", action="store_true", help="also seed the kb_docs knowledge base")
    args = parser.parse_args()

    db = get_db()
    print(f"Connected to {settings.mongodb_db}")

    counts = load_audit(db)
    for name, n in counts.items():
        print(f"  inserted {n:>4} into {name}")

    if args.kb:
        n = load_kb(db)
        print(f"  inserted {n:>4} into {settings.C_KB}")

    print("Seed complete.")


if __name__ == "__main__":
    main()
